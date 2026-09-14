"""
Ingestão: Pauta do Plenário e Pauta das Turmas — Corte Aberta
Fonte: transparencia.stf.jus.br/extensions/pauta_plenario e pauta_turmas

Única fonte prospectiva do site: processos já liberados para julgamento
colegiado, ainda não julgados — não é histórico de decisão. Mesmo
mecanismo de export nativo (.xlsx) de fetch_controle_concentrado.py.

Duas tabelas, porque as colunas disponíveis diferem entre os dois
painéis (Turmas não tem "RG Reconhecida" nem "Criminal"; usa "Orgão
Julgador" em vez de implícito "Plenário").

Execução:
  python3 ingestao/stf/fetch_pauta.py [--dry-run]
"""

import argparse
import io
import os
import re
import unicodedata
import warnings
from datetime import datetime

import openpyxl
from playwright.sync_api import sync_playwright
from supabase import create_client

from _corte_aberta_qlik import baixar_xlsx, bool_sim_nao, valor, NI

warnings.filterwarnings("ignore")

SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]

URL_PLENARIO = "https://transparencia.stf.jus.br/extensions/pauta_plenario/pauta_plenario.html"
URL_TURMAS = "https://transparencia.stf.jus.br/extensions/pauta_turmas/pauta_turmas.html"


def normalizar(txt: str) -> str:
    txt = unicodedata.normalize("NFD", txt)
    return "".join(c for c in txt if unicodedata.category(c) != "Mn").upper().strip()


def resolver_ministro(relator_bruto: str, ministros: list[dict]) -> str | None:
    if not relator_bruto or relator_bruto == NI:
        return None
    nome = re.sub(r"^MIN\.?\s+", "", relator_bruto.strip(), flags=re.IGNORECASE)
    nome_norm = normalizar(nome)
    for m in ministros:
        m_norm = normalizar(m["nome"])
        partes = m_norm.split()
        sobrenome = " ".join(partes[-2:]) if len(partes) > 1 else m_norm
        if sobrenome in nome_norm or m_norm in nome_norm or nome_norm in m_norm:
            return m["id"]
    return None


def data_iso(v):
    v = valor(v)
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    try:
        d, m, a = str(v).split("/")
        return f"{a}-{m}-{d}"
    except ValueError:
        return None


def processar_plenario(xlsx_bytes: bytes, ministros: list[dict]) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]

    linhas = {}
    for row in rows:
        r = dict(zip(header, row))
        classe, numero = valor(r.get("Classe")), valor(r.get("Numero"))
        if not classe or numero is None:
            continue
        relator_bruto = valor(r.get("Relator Atual"))
        vista_bruto = valor(r.get("Ministro Vista"))
        linhas[(classe, numero)] = {
            "classe": classe,
            "numero": numero,
            "ministro_id": resolver_ministro(relator_bruto, ministros),
            "relator_atual_bruto": relator_bruto,
            "data_autuacao": data_iso(r.get("Data Autuação")),
            "ramo_direito": valor(r.get("Ramo do Direito")),
            "criminal": bool_sim_nao(r.get("Criminal")),
            "sessao": valor(r.get("Sessão")),
            "rg_reconhecida": bool_sim_nao(r.get("RG Reconhecida")),
            "pedido_vista": bool_sim_nao(r.get("Pedido Vista")),
            "ministro_vista_id": resolver_ministro(vista_bruto, ministros),
            "ministro_vista_bruto": vista_bruto,
            "data_vista": data_iso(r.get("Data Vista")),
            "suspenso": bool_sim_nao(r.get("Suspenso")),
            "data_pauta": data_iso(r.get("Data Pauta")),
        }
    return list(linhas.values())


def processar_turmas(xlsx_bytes: bytes, ministros: list[dict]) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]

    linhas = {}
    for row in rows:
        r = dict(zip(header, row))
        classe, numero = valor(r.get("Classe")), valor(r.get("Numero"))
        orgao = valor(r.get("Orgão Julgador"))
        if not classe or numero is None or not orgao:
            continue
        relator_bruto = valor(r.get("Relator Atual"))
        vista_bruto = valor(r.get("Ministro Vista"))
        linhas[(classe, numero, orgao)] = {
            "classe": classe,
            "numero": numero,
            "orgao_julgador": orgao,
            "ministro_id": resolver_ministro(relator_bruto, ministros),
            "relator_atual_bruto": relator_bruto,
            "data_autuacao": data_iso(r.get("Data Autuação")),
            "ramo_direito": valor(r.get("Ramo do Direito")),
            "sessao": valor(r.get("Sessão")),
            "pedido_vista": bool_sim_nao(r.get("Pedido Vista")),
            "ministro_vista_id": resolver_ministro(vista_bruto, ministros),
            "ministro_vista_bruto": vista_bruto,
            "data_vista": data_iso(r.get("Data Vista")),
            "suspenso": bool_sim_nao(r.get("Suspenso")),
            "data_pauta": data_iso(r.get("Data Pauta")),
        }
    return list(linhas.values())


def run(dry_run: bool = False):
    sb = create_client(SUPABASE_URL, SUPABASE_KEY)
    ministros = sb.table("stf_ministros").select("id, nome").execute().data

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()

        print("Abrindo painel de Pauta do Plenário (navegador headless)...")
        xlsx_plenario = baixar_xlsx(page, URL_PLENARIO, "Relação de processos")
        lote_plenario = processar_plenario(xlsx_plenario, ministros)
        print(f"  {len(lote_plenario)} processos na pauta do plenário")

        print("Abrindo painel de Pauta das Turmas (navegador headless)...")
        xlsx_turmas = baixar_xlsx(page, URL_TURMAS, "Relação de processos")
        lote_turmas = processar_turmas(xlsx_turmas, ministros)
        print(f"  {len(lote_turmas)} processos na pauta das turmas")

        browser.close()

    if not dry_run:
        if lote_plenario:
            sb.table("stf_pauta_plenario").upsert(lote_plenario, on_conflict="classe,numero").execute()
        if lote_turmas:
            sb.table("stf_pauta_turmas").upsert(lote_turmas, on_conflict="classe,numero,orgao_julgador").execute()

    print(f"\n✅ {len(lote_plenario)} + {len(lote_turmas)} processos de pauta "
          f"{'processados (dry-run, nada salvo)' if dry_run else 'gravados'}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    run(args.dry_run)
