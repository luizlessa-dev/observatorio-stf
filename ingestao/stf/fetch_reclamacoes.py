"""
Ingestão: Reclamações constitucionais — Corte Aberta
Fonte: transparencia.stf.jus.br/extensions/reclamacoes (Qlik Sense)

Mesmo mecanismo de export nativo (.xlsx) de fetch_controle_concentrado.py
— ver ingestao/stf/_corte_aberta_qlik.py para a técnica de captura.

Bruto-primeiro: nenhum campo é reinterpretado, exceto a resolução best-
-effort de `ministro_id` a partir do nome do relator.

Execução:
  python3 ingestao/stf/fetch_reclamacoes.py [--dry-run]
"""

import argparse
import io
import os
import re
import unicodedata
import warnings
from datetime import datetime

import openpyxl
from playwright.sync_api import sync_playwright
from supabase import create_client

from _corte_aberta_qlik import baixar_xlsx, bool_sim_nao, valor, NI

warnings.filterwarnings("ignore")

SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]

URL_PAINEL = "https://transparencia.stf.jus.br/extensions/reclamacoes/reclamacoes.html"


def normalizar(txt: str) -> str:
    txt = unicodedata.normalize("NFD", txt)
    return "".join(c for c in txt if unicodedata.category(c) != "Mn").upper().strip()


def resolver_ministro(relator_bruto: str, ministros: list[dict]) -> str | None:
    if not relator_bruto or relator_bruto == NI:
        return None
    nome = re.sub(r"^MIN\.?\s+", "", relator_bruto.strip(), flags=re.IGNORECASE)
    nome_norm = normalizar(nome)
    for m in ministros:
        m_norm = normalizar(m["nome"])
        partes = m_norm.split()
        sobrenome = " ".join(partes[-2:]) if len(partes) > 1 else m_norm
        if sobrenome in nome_norm or m_norm in nome_norm or nome_norm in m_norm:
            return m["id"]
    return None


def data_iso(v):
    v = valor(v)
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    try:
        d, m, a = str(v).split("/")
        return f"{a}-{m}-{d}"
    except ValueError:
        return None


def processar(xlsx_bytes: bytes, ministros: list[dict]) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]

    linhas = {}
    for row in rows:
        r = dict(zip(header, row))
        processo = valor(r.get("Processo"))
        if not processo:
            continue
        relator_bruto = valor(r.get("Relator Atual"))
        linhas[processo] = {
            "processo": processo,
            "numero_unico": valor(r.get("Número Único")),
            "num_processos_origens": valor(r.get("Num. Processos Origens")),
            "data_autuacao": data_iso(r.get("Data Autuação")),
            "ministro_id": resolver_ministro(relator_bruto, ministros),
            "relator_atual_bruto": relator_bruto,
            "procedencia": valor(r.get("Procedência")),
            "preferencia_criminal": bool_sim_nao(r.get("Preferência Criminal")),
            "ramo_direito": valor(r.get("Ramo Direito")),
            "em_tramitacao": bool_sim_nao(r.get("Em Tramitação?")),
            "liminar_pendente": bool_sim_nao(r.get("Liminar Pendente")),
        }
    return list(linhas.values())


def run(dry_run: bool = False):
    sb = create_client(SUPABASE_URL, SUPABASE_KEY)
    ministros = sb.table("stf_ministros").select("id, nome").execute().data

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        print("Abrindo painel de Reclamações (navegador headless)...")
        xlsx_bytes = baixar_xlsx(page, URL_PAINEL, "Lista de Processos")
        browser.close()

    lote = processar(xlsx_bytes, ministros)
    resolvidos = sum(1 for l in lote if l["ministro_id"])
    print(f"  {len(lote)} reclamações ({resolvidos} com ministro_id resolvido)")

    if not dry_run:
        for i in range(0, len(lote), 500):
            sb.table("stf_reclamacoes").upsert(
                lote[i:i + 500], on_conflict="processo"
            ).execute()

    print(f"\n✅ {len(lote)} reclamações "
          f"{'processadas (dry-run, nada salvo)' if dry_run else 'gravadas'}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    run(args.dry_run)
