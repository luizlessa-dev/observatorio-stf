"""
Ingestão: Casos de Omissão Inconstitucional — Corte Aberta
Fonte: transparencia.stf.jus.br/extensions/omissao_inconstitucional

Lista curada (172 julgados) de casos em que o STF reconheceu omissão do
poder público — não um recorte de milhares de processos como os outros
painéis Corte Aberta. Mesmo mecanismo de export nativo (.xlsx).

Execução:
  python3 ingestao/stf/fetch_omissao_inconstitucional.py [--dry-run]
"""

import argparse
import io
import os
import re
import unicodedata
import warnings
from datetime import datetime

import openpyxl
from playwright.sync_api import sync_playwright
from supabase import create_client

from _corte_aberta_qlik import baixar_xlsx, valor, NI

warnings.filterwarnings("ignore")

SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]

URL_PAINEL = "https://transparencia.stf.jus.br/extensions/omissao_inconstitucional/omissao_inconstitucional.html"


def normalizar(txt: str) -> str:
    txt = unicodedata.normalize("NFD", txt)
    return "".join(c for c in txt if unicodedata.category(c) != "Mn").upper().strip()


def resolver_ministro(relator_bruto: str, ministros: list[dict]) -> str | None:
    if not relator_bruto or relator_bruto == NI:
        return None
    nome = re.sub(r"^MIN\.?\s+", "", relator_bruto.strip(), flags=re.IGNORECASE)
    nome_norm = normalizar(nome)
    for m in ministros:
        m_norm = normalizar(m["nome"])
        partes = m_norm.split()
        sobrenome = " ".join(partes[-2:]) if len(partes) > 1 else m_norm
        if sobrenome in nome_norm or m_norm in nome_norm or nome_norm in m_norm:
            return m["id"]
    return None


def limpar_multivalor(v):
    """'Administrativo;#1;#Saúde;#23' -> 'Administrativo, Saúde' — a fonte
    exporta campo multivalorado do Qlik como pares rótulo;#id-interno;
    o id não tem uso fora do próprio app."""
    v = valor(v)
    if v is None:
        return None
    partes = [p.strip() for p in v.split(";#") if p.strip() and not p.strip().isdigit()]
    return ", ".join(partes) if partes else None


def data_iso(v):
    v = valor(v)
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    try:
        d, m, a = str(v).split("/")
        return f"{a}-{m}-{d}"
    except ValueError:
        return None


def processar(xlsx_bytes: bytes, ministros: list[dict]) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]

    linhas = {}
    for row in rows:
        r = dict(zip(header, row))
        processo = valor(r.get("Processo"))
        if not processo:
            continue
        processo = processo.strip()
        relator_bruto = valor(r.get("Relator(a)"))
        linhas[processo] = {
            "materia": valor(r.get("Matéria")),
            "processo": processo,
            "ministro_id": resolver_ministro(relator_bruto, ministros),
            "relator_bruto": relator_bruto,
            "redator_acordao_bruto": valor(r.get("Redator(a) do acórdão")),
            "tipo_classe": valor(r.get("Tipo classe")),
            "classe_processo": valor(r.get("Classe processo")),
            "incidente": valor(r.get("Incidente")),
            "numero_processo": valor(r.get("Número processo")),
            "data_julgamento": data_iso(r.get("Data julgamento")),
            "ementa": valor(r.get("Texto ementa")),
            "orgao_julgador": valor(r.get("Órgão julgador")),
            "tipo_omissao": limpar_multivalor(r.get("Tipo de omissão")),
            "ramo_direito": limpar_multivalor(r.get("Ramo do Direito")),
            "link_processo": valor(r.get("Link processo")),
            "link_inteiro_teor": valor(r.get("Link inteiro teor acordão")),
            "link_jurisprudencia": valor(r.get("Link pesquisa jurisprudência")),
        }
    return list(linhas.values())


def run(dry_run: bool = False):
    sb = create_client(SUPABASE_URL, SUPABASE_KEY)
    ministros = sb.table("stf_ministros").select("id, nome").execute().data

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        print("Abrindo painel de Omissão Inconstitucional (navegador headless)...")
        xlsx_bytes = baixar_xlsx(page, URL_PAINEL, "Lista de dados")
        browser.close()

    lote = processar(xlsx_bytes, ministros)
    resolvidos = sum(1 for l in lote if l["ministro_id"])
    print(f"  {len(lote)} casos ({resolvidos} com ministro_id resolvido)")

    if not dry_run and lote:
        sb.table("stf_omissao_inconstitucional").upsert(lote, on_conflict="processo").execute()

    print(f"\n✅ {len(lote)} casos de omissão inconstitucional "
          f"{'processados (dry-run, nada salvo)' if dry_run else 'gravados'}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    run(args.dry_run)
