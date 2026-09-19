"""
Ingestão: Informação à Sociedade — Corte Aberta
Fonte: transparencia.stf.jus.br/extensions/Informacao_A_Sociedade (Qlik Sense)

Como Ações Covid-19, é uma curadoria do STF (178 julgados selecionados),
não um recorte automático — o projeto existe para explicar em linguagem
simples os julgamentos de maior repercussão. A aba padrão do painel
("Julgamentos") é um feed de cards, sem tabela exportável; a aba "Lista"
tem a tabela real com export nativo (.xlsx), usada aqui.

Execução:
  python3 ingestao/stf/fetch_informacao_sociedade.py [--dry-run]
"""

import argparse
import io
import os
import re
import unicodedata
import warnings
from datetime import datetime

import openpyxl
from playwright.sync_api import sync_playwright
from supabase import create_client

from _corte_aberta_qlik import baixar_xlsx, valor, NI

warnings.filterwarnings("ignore")

SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]

URL_PAINEL = "https://transparencia.stf.jus.br/extensions/Informacao_A_Sociedade/Informacao_A_Sociedade.html"


def normalizar(txt: str) -> str:
    # relator_bruto às vezes vem com quebra de linha embutida no meio do
    # nome (achado real: "Min. Alexandre de\nMoraes") — colapsa qualquer
    # sequência de espaço em branco antes de comparar.
    txt = unicodedata.normalize("NFD", txt)
    txt = "".join(c for c in txt if unicodedata.category(c) != "Mn").upper().strip()
    return re.sub(r"\s+", " ", txt)


def resolver_ministro(relator_bruto: str, ministros: list[dict]) -> str | None:
    """2 dos 178 julgados têm mais de um relator (ex. "Min. Flávio Dino e
    Min. Dias Toffoli") — resolve pro primeiro que bater; relator_bruto
    preserva o texto completo pra quem quiser ver os dois."""
    if not relator_bruto or relator_bruto == NI:
        return None
    nome = re.sub(r"^MIN\.?\s+", "", relator_bruto.strip(), flags=re.IGNORECASE)
    nome_norm = normalizar(nome)
    for m in ministros:
        m_norm = normalizar(m["nome"])
        partes = m_norm.split()
        sobrenome = " ".join(partes[-2:]) if len(partes) > 1 else m_norm
        if sobrenome in nome_norm or m_norm in nome_norm or nome_norm in m_norm:
            return m["id"]
    return None


def data_iso(v):
    v = valor(v)
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    try:
        d, m, a = str(v).split("/")
        return f"{a}-{m}-{d}"
    except ValueError:
        return None


def processar(dados: bytes, ministros: list[dict]) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]

    linhas = {}
    for row in rows:
        r = dict(zip(header, row))
        processo = valor(r.get("Processo"))
        data_julgamento = data_iso(r.get("Data Julgamento"))
        if not processo:
            continue
        relator_bruto = valor(r.get("Relator"))
        numero_bruto = r.get("Número")
        linhas[(processo, data_julgamento)] = {
            "processo": processo,
            "classe": valor(r.get("Classe")),
            "numero": str(numero_bruto) if numero_bruto is not None else None,
            "data_julgamento": data_julgamento,
            "ministro_id": resolver_ministro(relator_bruto, ministros),
            "relator_bruto": relator_bruto,
            "fatos": valor(r.get("Fatos")),
            "fundamentos_decisao": valor(r.get("Fundamentos da Decisão")),
            "questoes_juridicas": valor(r.get("Questões Jurídicas")),
            "tese": valor(r.get("Tese")),
            "resultado": valor(r.get("Resultado")),
            "placar": valor(r.get("Placar")),
            "voto_prevaleceu": valor(r.get("Voto que Prevaleceu")),
            "votos_divergentes": valor(r.get("Votos Divergentes")),
            "votacao": valor(r.get("Votação")),
            "ods": valor(r.get("ODS")),
            "ambiente_julgamento": valor(r.get("Ambiente de Julgamento")),
            "resumo_pt": valor(r.get("Resumo em Português")),
            "resumo_en": valor(r.get("Resumo em Inglês")),
            "resumo_es": valor(r.get("Resumo em Espanhol")),
        }
    return list(linhas.values())


def run(dry_run: bool = False, arquivo_local: str | None = None):
    sb = create_client(SUPABASE_URL, SUPABASE_KEY)
    ministros = sb.table("stf_ministros").select("id, nome").execute().data

    if arquivo_local:
        with open(arquivo_local, "rb") as f:
            dados = f.read()
    else:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            print("Abrindo painel de Informação à Sociedade (navegador headless)...")
            dados = baixar_xlsx(page, URL_PAINEL, "Lista")
            browser.close()

    lote = processar(dados, ministros)
    resolvidos = sum(1 for l in lote if l["ministro_id"])
    print(f"  {len(lote)} julgados ({resolvidos} com ministro_id resolvido, "
          f"{len(lote) - resolvidos} sem correspondência)")

    if not dry_run and lote:
        sb.table("stf_informacao_sociedade").upsert(lote, on_conflict="processo,data_julgamento").execute()

    print(f"\n✅ {len(lote)} julgados de Informação à Sociedade "
          f"{'processados (dry-run, nada salvo)' if dry_run else 'gravados'}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--arquivo-local", help="Usa um .xlsx já baixado em vez de acessar o painel")
    args = parser.parse_args()
    run(args.dry_run, args.arquivo_local)
