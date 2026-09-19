"""
Ingestão: Registro e Distribuição de processos — Corte Aberta
Fonte: transparencia.stf.jus.br/extensions/distribuidos (Qlik Sense)

Mesmo padrão de fetch_recebimento_baixa.py: export xlsx (zip, magic "PK"),
escopo do ANO CORRENTE via baixar_ano_atual() (reaproveitado do mesmo
módulo — evita repetir a função) para não estourar pra histórico
completo, um único download já traz os dois tipos de andamento
("Registrado à Presidência" e "Distribuído aos Ministros") combinados.

Diferente do painel anterior: aqui quase toda linha tem "Ministro(a)"
preenchido — é o próprio evento de distribuição. A exceção é o lado
"Registrado à Presidência", que na prática é quase todo rotulado
"MINISTRO PRESIDENTE" (cargo, não pessoa — fica sem ministro_id, por
desenho, não por falha de resolução).

Execução:
  python3 ingestao/stf/fetch_distribuicao.py [--dry-run]
"""

import argparse
import io
import os
import re
import unicodedata
import warnings
from datetime import datetime

import openpyxl
from playwright.sync_api import sync_playwright
from supabase import create_client

from fetch_recebimento_baixa import baixar_ano_atual

warnings.filterwarnings("ignore")

SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]

URL_PAINEL = "https://transparencia.stf.jus.br/extensions/distribuidos/distribuidos.html"


def normalizar(txt: str) -> str:
    txt = unicodedata.normalize("NFD", txt)
    return "".join(c for c in txt if unicodedata.category(c) != "Mn").upper().strip()


def resolver_ministro(ministro_bruto: str, ministros: list[dict]) -> str | None:
    if not ministro_bruto:
        return None
    nome = re.sub(r"^MIN\.?\s+", "", ministro_bruto.strip(), flags=re.IGNORECASE)
    nome_norm = normalizar(nome)
    for m in ministros:
        m_norm = normalizar(m["nome"])
        partes = m_norm.split()
        sobrenome = " ".join(partes[-2:]) if len(partes) > 1 else m_norm
        if sobrenome in nome_norm or m_norm in nome_norm or nome_norm in m_norm:
            return m["id"]
    return None


def valor(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return None if v in ("", "-", "NÃO INFORMADO") else v
    return v


def data_iso(v):
    v = valor(v)
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    try:
        d, m, a = str(v).split("/")
        return f"{a}-{m}-{d}"
    except ValueError:
        return None


def bool_sim_nao(v):
    v = valor(v)
    if v is None:
        return None
    return str(v).strip().lower() == "sim"


def processar(dados: bytes, ministros: list[dict]) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]

    linhas = []
    for row in rows:
        r = dict(zip(header, row))
        classe = valor(r.get("Classe"))
        numero_bruto = r.get("Nº do processo")
        if not classe or numero_bruto in (None, "", "-"):
            continue
        try:
            numero = int(numero_bruto)
        except (ValueError, TypeError):
            continue
        ministro_bruto = valor(r.get("Ministro(a)"))
        linhas.append({
            "tipo_andamento": valor(r.get("Tipo de andamento")),
            "classe": classe,
            "numero": numero,
            "ministro_id": resolver_ministro(ministro_bruto, ministros),
            "ministro_bruto": ministro_bruto,
            "link_processo": valor(r.get("Link")),
            "ultima_localizacao": valor(r.get("Última localização")),
            "data_autuacao": data_iso(r.get("Data da autuação")),
            "data_baixa": data_iso(r.get("Data da baixa")),
            "em_tramitacao": bool_sim_nao(r.get("Em tramitação")),
            "grupo_origem": valor(r.get("Grupo origem")),
            "meio_processo": valor(r.get("Meio processo")),
            "data_andamento": data_iso(r.get("Data do andamento")),
            "andamento": valor(r.get("Andamento")),
            "subgrupo_andamento": valor(r.get("Subgrupo do andamento")),
            "substituicao_redistribuicao": bool_sim_nao(r.get("Indicador de substituição ou redistribuição")),
            "orgao_origem": valor(r.get("Orgão origem")),
            "procedencia": valor(r.get("Procedência")),
            "ramo_direito": valor(r.get("Ramo do direito")),
            "assunto_completo": valor(r.get("Assunto completo")),
            "polo_ativo": valor(r.get("Polo ativo")),
            "advogado_polo_ativo": valor(r.get("Advogado polo ativo")),
            "polo_passivo": valor(r.get("Polo passivo")),
            "advogado_polo_passivo": valor(r.get("Advogado polo passivo")),
        })
    # 319 processos foram redistribuídos/registrados mais de uma vez no
    # mesmo ano — (classe, numero, tipo_andamento, data_andamento) resolve
    # quase tudo, mas sobram 2 pares com a mesma data também (confirmado
    # numa carga real, ver migration 0032). Dedup aqui evita que o upsert
    # em lote rejeite o lote inteiro por "affect row a second time".
    vistos = set()
    deduplicadas = []
    for l in linhas:
        chave = (l["classe"], l["numero"], l["tipo_andamento"], l["data_andamento"])
        if chave in vistos:
            continue
        vistos.add(chave)
        deduplicadas.append(l)
    return deduplicadas


def run(dry_run: bool = False, arquivo_local: str | None = None):
    sb = create_client(SUPABASE_URL, SUPABASE_KEY)
    ministros = sb.table("stf_ministros").select("id, nome").execute().data

    if arquivo_local:
        with open(arquivo_local, "rb") as f:
            dados = f.read()
    else:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            print("Abrindo painel de Registro e Distribuição (navegador headless)...")
            dados = baixar_ano_atual(page, URL_PAINEL, "Lista de processos")
            browser.close()

    lote = processar(dados, ministros)
    resolvidos = sum(1 for l in lote if l["ministro_id"])
    print(f"  {len(lote)} eventos ({resolvidos} com ministro_id resolvido, "
          f"{len(lote) - resolvidos} sem correspondência)")

    if not dry_run:
        for i in range(0, len(lote), 500):
            sb.table("stf_distribuicao").upsert(
                lote[i:i + 500], on_conflict="classe,numero,tipo_andamento,data_andamento"
            ).execute()

    print(f"\n✅ {len(lote)} eventos de registro/distribuição "
          f"{'processados (dry-run, nada salvo)' if dry_run else 'gravados'}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--arquivo-local", help="Usa um export já baixado em vez de acessar o painel")
    args = parser.parse_args()
    run(args.dry_run, args.arquivo_local)
